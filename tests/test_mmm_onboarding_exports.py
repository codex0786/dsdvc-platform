import datetime as dt
import io
import os

import numpy as np
import pandas as pd
import pytest

from dsdvc.mmm import analysis as A
from dsdvc.mmm.model import LiftTest, MarketingMixModel, MMMSpec, geometric_adstock, hill
from dsdvc.mmm.synthetic import TRUTH, messy_client_export, synthetic_mmm_data
from dsdvc.onboarding.mapping import _parse_number, apply_mapping, parse_dates, suggest_mapping
from dsdvc.onboarding.templates import MARKETING_WEEKLY, template_frame
from dsdvc.onboarding.validate import validate

SPEC = MMMSpec(target_col="revenue", channels=list(TRUTH), controls=["price_index", "promo"])


@pytest.fixture(scope="module")
def data():
    return synthetic_mmm_data(156)


@pytest.fixture(scope="module")
def fit(data):
    return MarketingMixModel(SPEC).fit(data[0])


def test_adstock_and_hill():
    x = np.array([1.0, 0, 0, 0])
    a = geometric_adstock(x, 0.5)
    np.testing.assert_allclose(a, [0.5, 0.25, 0.125, 0.0625])
    assert geometric_adstock(np.ones(500), 0.7)[-1] == pytest.approx(1.0)  # normalised: steady state = spend
    assert hill(np.array([1.0]), 1.0, 2.0)[0] == pytest.approx(0.5)


def test_mmm_recovers_roi_of_well_identified_channels(fit, data):
    truth = data[1]["roi"]
    s = A.channel_summary(fit).set_index("channel")
    assert fit.diagnostics["converged"] and fit.diagnostics["r2"] > 0.85
    for c in ("meta_ads", "google_search", "influencers"):
        assert s.loc[c, "roi_low"] <= truth[c] * 1.1 and truth[c] * 0.9 <= s.loc[c, "roi_high"], c
    order = s.roi.sort_values().index
    assert order[-1] == "influencers" and order[0] == "ooh"


def test_lift_test_calibration_moves_estimate_toward_truth(data):
    df, truth = data
    th, K, S, b = TRUTH["youtube"]
    x = df.youtube.to_numpy() / df.youtube[df.youtube > 0].mean()
    xz = x.copy()
    xz[100:108] = 0
    lift = float((b * hill(geometric_adstock(x, th), K, S) - b * hill(geometric_adstock(xz, th), K, S))[100:].sum() * 6_000_000)
    lt = LiftTest("youtube", df.date[100], df.date[107], lift, 0.15 * lift)
    base = A.channel_summary(MarketingMixModel(SPEC).fit(df)).set_index("channel").loc["youtube", "roi"]
    cal = A.channel_summary(MarketingMixModel(SPEC).fit(df, [lt])).set_index("channel").loc["youtube", "roi"]
    assert abs(cal - truth["roi"]["youtube"]) < abs(base - truth["roi"]["youtube"])


def test_plan_respects_budget_and_bounds(fit):
    p = A.optimize_budget(fit)
    t = p.table
    assert t.recommended_weekly.sum() == pytest.approx(p.weekly_budget, rel=1e-4)
    assert (t.recommended_weekly >= t.lower_bound - 1e-6).all() and (t.recommended_weekly <= t.upper_bound + 1e-6).all()
    assert p.uplift_low <= p.uplift <= p.uplift_high and 0 <= p.p_better <= 1
    prof = A.optimize_budget(fit, objective="profit", margin=0.3)
    assert prof.table.recommended_weekly.sum() <= prof.weekly_budget * (1 + 1e-6)


def test_backtest_passes_on_planted_data(data):
    bt = A.backtest(MarketingMixModel(SPEC), data[0], 12)
    assert bt.metrics["mape"] < 0.1 and bt.verdict.startswith("PASS")
    assert bt.series.holdout.sum() == 12


def test_response_curves_are_monotone(fit):
    rc = A.response_curves(fit)
    for _, g in rc.groupby("channel"):
        assert (np.diff(g.weekly_effect.to_numpy()) >= -1e-9).all()


def test_identifiability_warnings_for_flat_and_collinear_spend(data):
    df = data[0].copy()
    df["flat"] = 100000.0
    df["copy_meta"] = df.meta_ads * 1.01
    spec = MMMSpec(target_col="revenue", channels=["meta_ads", "copy_meta", "flat"], controls=[])
    w = " ".join(MarketingMixModel(spec, n_restarts=2).fit(df).diagnostics["warnings"])
    assert "barely varies" in w and "move together" in w


def test_number_and_date_parsing():
    vals = ["₹1,20,000", "2.5 Cr", "12 lakh", "(5,000)", "12%", "Yes", "no", "abc", "", None]
    assert [_parse_number(v) for v in vals] == [120000.0, 25000000.0, 1200000.0, -5000.0, 0.12, 1.0, 0.0, None, None, None]
    d, how = parse_dates(pd.Series(["13/01/2025", "20/01/2025", "27/01/2025"]))
    assert how == "day-first" and d.iloc[0] == pd.Timestamp("2025-01-13")


def test_mapping_round_trips_messy_export(data):
    raw = messy_client_export()
    m = suggest_mapping(raw, MARKETING_WEEKLY)
    mm = m.as_dict()
    assert mm["FB + IG Spend (₹)"] == "spend_meta_ads" and mm["Total Sales (INR)"] == "revenue" and "Notes" not in mm
    canon, notes = apply_mapping(raw, m, MARKETING_WEEKLY)
    df = data[0]
    assert (canon.revenue.values == df.revenue.values).all() and (canon.spend_youtube.values == df.youtube.values).all()
    assert validate(canon, MARKETING_WEEKLY).status == "READY"


def test_long_format_is_pivoted():
    long = pd.DataFrame({"Week": ["06/01/2025"] * 2 + ["13/01/2025"] * 2, "Channel": ["Facebook", "Google Ads"] * 2,
                         "Spend": ["1,000", "2,000", "1,500", "2,500"], "Sales": ["10,000", "10,000", "12,000", "12,000"]})
    m = suggest_mapping(long, MARKETING_WEEKLY)
    canon, notes = apply_mapping(long, m, MARKETING_WEEKLY)
    assert {"spend_meta_ads", "spend_google_search"} <= set(canon.columns) and len(canon) == 2
    assert any("long format" in n for n in notes)


def test_validation_blocks_bad_data():
    df = template_frame("marketing_weekly", 3)
    df["date"] = pd.to_datetime(df["date"])
    df.loc[1, "spend_meta_ads"] = -5
    sc = validate(df, MARKETING_WEEKLY)
    assert sc.status == "BLOCKED"
    names = {c.name for c in sc.checks if c.status == "fail"}
    assert "negative spend_meta_ads" in names and "history length" in names


def test_workspace_requires_attestation(tmp_path, monkeypatch):
    from dsdvc.database.models import SessionLocal, init_db
    from dsdvc.onboarding.workspace import load_dataset, save_dataset
    monkeypatch.setenv("DSDVC_DATA_DIR", str(tmp_path))
    init_db()
    raw = messy_client_export()
    m = suggest_mapping(raw, MARKETING_WEEKLY)
    canon, _ = apply_mapping(raw, m, MARKETING_WEEKLY)
    sc = validate(canon, MARKETING_WEEKLY)
    with SessionLocal() as db:
        with pytest.raises(ValueError):
            save_dataset(db, df=canon, dataset="marketing_weekly", filename="x.csv", mapping=m, scorecard=sc, attestation=" ")
        row = save_dataset(db, df=canon, dataset="marketing_weekly", filename="x.csv", mapping=m, scorecard=sc, attestation="test")
        back = load_dataset(row)
    assert len(back) == 156 and row.sha256 and row.status == "READY"


def test_exports_open_and_contain_plan(fit, data):
    from openpyxl import load_workbook
    from pptx import Presentation
    from dsdvc.exports.deck import display_name, executive_deck
    from dsdvc.exports.excel import media_plan_workbook
    bt = A.backtest(MarketingMixModel(SPEC, n_restarts=2), data[0], 8)
    plan, summ = A.optimize_budget(fit), A.channel_summary(fit)
    prov = {"filename": "f.csv", "rows": 156, "sha256": "abc", "validation": "READY", "attestation": "t", "data_version": "v"}
    x = media_plan_workbook(client="Acme", plan=plan, summary=summ, curves=A.response_curves(fit), backtest=bt,
                            params=A.channel_params(fit), horizon_weeks=13, start=dt.date(2026, 1, 5),
                            diagnostics=fit.diagnostics, provenance=prov)
    wb = load_workbook(io.BytesIO(x))
    assert wb.sheetnames == ["Summary", "Channel plan", "Weekly schedule", "Channel ROI", "Response curves", "Model accuracy",
                             "Data & assumptions"]
    sched = wb["Weekly schedule"]
    assert sched.max_row == 4 + 13
    d = executive_deck(client="Acme", plan=plan, summary=summ, decomposition=A.decomposition(fit), backtest=bt,
                       diagnostics=fit.diagnostics, provenance=prov, horizon_weeks=13)
    prs = Presentation(io.BytesIO(d))
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert len(prs.slides) == 9 and "Acme" in text and "will go viral" not in text.lower()
    assert display_name("spend_google_search") == "Google Search" and display_name("ooh") == "OOH"


def test_api_onboarding_to_export():
    from fastapi.testclient import TestClient
    from dsdvc.api.main import app
    H = {"X-API-Key": os.environ["DSDVC_BOOTSTRAP_ADMIN_KEY"]}
    csv = messy_client_export().to_csv(index=False).encode()
    with TestClient(app) as c:
        assert c.get("/onboarding/templates/marketing_weekly", headers=H).json()["key"] == "marketing_weekly"
        pv = c.post("/onboarding/preview", headers=H, files={"file": ("client.csv", csv, "text/csv")}).json()
        assert pv["status"] == "READY"
        assert c.post("/onboarding/datasets", headers=H, files={"file": ("client.csv", csv, "text/csv")},
                      data={"attestation": " "}).status_code == 422
        ds = c.post("/onboarding/datasets", headers=H, files={"file": ("client.csv", csv, "text/csv")},
                    data={"attestation": "test fixture, no personal data"}).json()
        fitr = c.post("/mmm/fit", headers=H, json={"dataset_id": ds["dataset_id"], "holdout_weeks": 8}).json()
        assert fitr["backtest"]["verdict"] and fitr["model_version"]
        rid = fitr["run_id"]
        assert c.get(f"/mmm/{rid}/export/xlsx", headers=H).status_code == 409
        plan = c.post(f"/mmm/{rid}/plan", headers=H, json={}).json()
        assert plan["status"] == "RECOMMENDED"
        x = c.get(f"/mmm/{rid}/export/xlsx", headers=H, params={"client": "Acme"})
        p = c.get(f"/mmm/{rid}/export/pptx", headers=H, params={"client": "Acme"})
        assert x.status_code == 200 and x.content[:2] == b"PK" and p.status_code == 200 and p.content[:2] == b"PK"
