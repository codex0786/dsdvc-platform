"""Client-ready Excel media plan (openpyxl) with native charts.

Sheets: Summary · Channel plan · Weekly schedule · Channel ROI · Response curves · Model accuracy ·
Data & assumptions. Every sheet states the epistemic status of its numbers.
"""
from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dsdvc import MODEL_VERSION, __version__
from dsdvc.exports.deck import display_name

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True, color="1F3A5F")
NOTE_FONT = Font(italic=True, color="52514E", size=9)
THIN = Side(style="thin", color="D5D4CF")
CURRENCY = '#,##0'
PCT = '0.0%'
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]


def _table(ws, df: pd.DataFrame, start_row: int, start_col: int = 1, formats: dict | None = None) -> tuple[int, int]:
    formats = formats or {}
    for j, col in enumerate(df.columns):
        c = ws.cell(start_row, start_col + j, str(col))
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime().date()
            c = ws.cell(start_row + i, start_col + j, v)
            c.border = Border(bottom=THIN)
            fmt = formats.get(df.columns[j])
            if fmt:
                c.number_format = fmt
    for j, col in enumerate(df.columns):
        width = max(len(str(col)), *(len(f"{v:,.0f}" if isinstance(v, (int, float)) else str(v)) for v in df[col].head(50))) + 2
        ws.column_dimensions[get_column_letter(start_col + j)].width = min(max(width, 10), 42)
    return start_row + len(df), start_col + len(df.columns) - 1


def _title(ws, text: str, status: str, note: str = ""):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Status: {status}" + (f" · {note}" if note else "")
    ws["A2"].font = NOTE_FONT


def media_plan_workbook(*, client: str, plan, summary: pd.DataFrame, curves: pd.DataFrame, backtest, params: pd.DataFrame,
                        horizon_weeks: int, start: date, diagnostics: dict, provenance: dict, decision: str = "",
                        currency: str = "INR") -> bytes:
    wb = Workbook()
    # ---------------- Summary
    ws = wb.active
    ws.title = "Summary"
    _title(ws, f"{client} — media plan", "RECOMMENDED", "built from an ESTIMATED marketing mix model")
    t = plan.table
    rows = [
        ("Decision", decision or "Reallocate the weekly media budget to maximise incremental outcome"),
        ("Prepared", date.today().isoformat()),
        ("Planning horizon", f"{horizon_weeks} weeks from {start.isoformat()}"),
        ("Weekly budget", plan.weekly_budget), ("Total budget over horizon", plan.weekly_budget * horizon_weeks),
        ("Expected incremental outcome / week vs current mix", plan.uplift),
        ("  90% interval (low)", plan.uplift_low), ("  90% interval (high)", plan.uplift_high),
        ("Probability the plan beats the current mix", plan.p_better),
        ("Expected incremental outcome over horizon", plan.uplift * horizon_weeks),
        ("Model accuracy (back-test)", backtest.verdict if backtest else "not run"),
        ("Back-test MAPE", backtest.metrics["mape"] if backtest else None),
        ("Currency", currency), ("Model version", MODEL_VERSION), ("Software", f"dsdvc {__version__}"),
        ("Data version", provenance.get("data_version", "")),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(i, 1, k).font = Font(bold=True)
        c = ws.cell(i, 2, v)
        if isinstance(v, float):
            c.number_format = PCT if "Probability" in k or "MAPE" in k else CURRENCY
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 52, 70
    r = 4 + len(rows) + 1
    ws.cell(r, 1, "Notes").font = Font(bold=True)
    for k, n in enumerate(plan.notes + diagnostics.get("warnings", []), start=1):
        ws.cell(r + k, 1, f"• {n}").alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r + k, start_column=1, end_row=r + k, end_column=2)
        ws.row_dimensions[r + k].height = 30

    # ---------------- Channel plan
    ws = wb.create_sheet("Channel plan")
    _title(ws, "Channel plan (weekly)", "RECOMMENDED", "current = average weekly spend in the modelled history")
    t = t.assign(channel=t.channel.map(display_name))
    cp = t[["channel", "current_weekly", "recommended_weekly", "change_pct", "current_weekly_effect", "recommended_weekly_effect",
            "lower_bound", "upper_bound"]].copy()
    cp["horizon_total"] = cp.recommended_weekly * horizon_weeks
    cp.columns = ["Channel", "Current / week", "Recommended / week", "Change", "Current effect / week",
                  "Recommended effect / week", "Min allowed / week", "Max allowed / week", f"Total over {horizon_weeks} wks"]
    end, _ = _table(ws, cp, 4, formats={c: CURRENCY for c in cp.columns[1:] if c != "Change"} | {"Change": PCT})
    tot = end + 1
    ws.cell(tot, 1, "Total").font = Font(bold=True)
    for j in (2, 3, 5, 6, 9):
        L = get_column_letter(j)
        c = ws.cell(tot, j, f"=SUM({L}5:{L}{end})")
        c.number_format, c.font = CURRENCY, Font(bold=True)
    ws.freeze_panes = "B5"

    # ---------------- Weekly schedule
    ws = wb.create_sheet("Weekly schedule")
    _title(ws, "Weekly schedule", "RECOMMENDED", "flat weekly allocation; adjust for flighting/seasonality with the planner")
    weeks = pd.date_range(pd.Timestamp(start), periods=horizon_weeks, freq="7D")
    sched = pd.DataFrame({"Week starting": weeks})
    for ch, v in zip(t.channel, t.recommended_weekly):
        sched[ch] = round(v, -2)
    sched["Total"] = sched[list(t.channel)].sum(axis=1)
    _table(ws, sched, 4, formats={c: CURRENCY for c in sched.columns[1:]} | {"Week starting": "yyyy-mm-dd"})
    ws.freeze_panes = "B5"

    # ---------------- Channel ROI
    ws = wb.create_sheet("Channel ROI")
    _title(ws, "Return by channel (modelled history)", "ESTIMATED", "90% intervals from the model posterior")
    roi = summary.assign(channel=summary.channel.map(display_name))[["channel", "spend", "contribution", "contribution_low", "contribution_high", "roi", "roi_low", "roi_high",
                   "mroi", "p_roi_above_1", "share_of_spend", "share_of_effect"]].copy()
    roi.columns = ["Channel", "Spend", "Incremental outcome", "low", "high", "ROI", "ROI low", "ROI high",
                   "Marginal ROI", "P(ROI>1)", "Share of spend", "Share of effect"]
    _table(ws, roi, 4, formats={"Spend": CURRENCY, "Incremental outcome": CURRENCY, "low": CURRENCY, "high": CURRENCY,
                                "ROI": "0.00", "ROI low": "0.00", "ROI high": "0.00", "Marginal ROI": "0.00",
                                "P(ROI>1)": PCT, "Share of spend": PCT, "Share of effect": PCT})
    ws.cell(5 + len(roi) + 1, 1, "ROI = incremental outcome ÷ spend. Marginal ROI = return on the next unit of spend at "
            "current levels — the number that should drive reallocation.").font = NOTE_FONT
    pr = params.assign(channel=params.channel.map(display_name))
    _table(ws, pr.round(3), 8 + len(roi), formats={"half_saturation_spend": CURRENCY})

    # ---------------- Response curves
    ws = wb.create_sheet("Response curves")
    _title(ws, "Response curves (weekly, steady state)", "ESTIMATED", "incremental outcome per week at a sustained weekly spend")
    wide = curves.pivot_table(index="weekly_spend", columns="channel", values="weekly_effect").reset_index()
    col = 1
    chart = ScatterChart()
    chart.title, chart.x_axis.title, chart.y_axis.title = "Weekly effect vs weekly spend", "Weekly spend", "Weekly effect"
    chart.height, chart.width = 10, 20
    for k, ch_raw in enumerate(curves.channel.unique()):
        ch = display_name(ch_raw)
        sub = curves[curves.channel == ch_raw][["weekly_spend", "weekly_effect", "low", "high"]].reset_index(drop=True)
        sub.columns = [f"{ch} spend", f"{ch} effect", "low", "high"]
        end, _ = _table(ws, sub, 4, col, formats={c: CURRENCY for c in sub.columns})
        xs = Reference(ws, min_col=col, min_row=5, max_row=end)
        ys = Reference(ws, min_col=col + 1, min_row=4, max_row=end)
        s = Series(ys, xs, title_from_data=True)
        s.graphicalProperties.line.solidFill = SERIES[k % len(SERIES)]
        s.graphicalProperties.line.width = 22000
        s.marker.symbol = "none"
        s.smooth = True
        chart.series.append(s)
        col += 5
    ws.add_chart(chart, f"A{len(wide) + 8}")

    # ---------------- Model accuracy
    ws = wb.create_sheet("Model accuracy")
    if backtest:
        _title(ws, "Model accuracy (back-test)", "ESTIMATED", backtest.verdict)
        m = backtest.metrics
        kv = [("Holdout weeks", m["holdout_weeks"]), ("MAPE (holdout)", m["mape"]), ("Total error (holdout)", m["total_error_pct"]),
              ("R² (holdout)", m["r2"]), ("90% interval coverage", m["interval_coverage_90"]), ("R² (training)", m["train_r2"])]
        for i, (k, v) in enumerate(kv, start=4):
            ws.cell(i, 1, k).font = Font(bold=True)
            c = ws.cell(i, 2, v)
            c.number_format = PCT if ("MAPE" in k or "error" in k or "coverage" in k) else "0.00"
        ser = backtest.series[["date", "actual", "predicted", "low", "high", "holdout"]].copy()
        ser["holdout"] = ser.holdout.map({True: "holdout", False: "train"})
        end, _ = _table(ws, ser, 12, formats={"actual": CURRENCY, "predicted": CURRENCY, "low": CURRENCY, "high": CURRENCY,
                                              "date": "yyyy-mm-dd"})
        lc = LineChart()
        lc.title, lc.height, lc.width = "Actual vs predicted (last weeks = unseen holdout)", 9, 22
        lc.add_data(Reference(ws, min_col=2, min_row=12, max_col=3, max_row=end), titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=1, min_row=13, max_row=end))
        for s, colr in zip(lc.series, ("52514E", "2A78D6")):
            s.graphicalProperties.line.solidFill = colr
            s.graphicalProperties.line.width = 20000
        ws.add_chart(lc, "H4")
    else:
        _title(ws, "Model accuracy", "NOT RUN", "run a back-test before sharing recommendations")

    # ---------------- Data & assumptions
    ws = wb.create_sheet("Data & assumptions")
    _title(ws, "Data, method and limitations", "DOCUMENTATION")
    items = [("Dataset", provenance.get("filename", "")), ("Rows (weeks)", provenance.get("rows", "")),
             ("SHA-256 (canonical data)", provenance.get("sha256", "")), ("Validation status", provenance.get("validation", "")),
             ("Authorisation", provenance.get("attestation", "")), ("Uploaded by", provenance.get("uploaded_by", "")),
             ("Model", "Bayesian marketing mix model: geometric adstock + Hill saturation, trend, yearly seasonality, controls; "
                       "MAP + Laplace posterior; calibrated with lift tests where supplied"),
             ("Lift tests used", provenance.get("lift_tests", 0)), ("Model fit R²", diagnostics.get("r2")),
             ("Converged", diagnostics.get("converged")),
             ("Limitations", "Correlational model of historical variation. Plans outside the historical spend range are "
                             "extrapolations; validate large shifts with geo-lift tests. Results assume no structural change "
                             "(new competitors, pricing, distribution) during the horizon.")]
    for i, (k, v) in enumerate(items, start=4):
        ws.cell(i, 1, k).font = Font(bold=True)
        c = ws.cell(i, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
